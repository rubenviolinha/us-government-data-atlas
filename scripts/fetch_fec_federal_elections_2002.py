#!/usr/bin/env python3
"""Fetch the FEC's 2002 Senate and House Excel publications (no API key)."""

from datetime import datetime, timezone
import hashlib
import io
import json
from pathlib import Path
from urllib.request import Request, urlopen

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
RAW = ROOT / "data" / "raw"
NORMALIZED = ROOT / "data" / "normalized"
SOURCES = {
    "senate": "https://www.fec.gov/documents/1740/FederalElections2002_Senate.xlsx",
    "house": "https://www.fec.gov/documents/1741/FederalElections2002_House.xlsx",
}


def clean(value):
    return str(value).strip() if value is not None else ""


def synthetic_id(year, office, state, district, name, party):
    key = "|".join((str(year), office, state, district, name, party))
    return "LEGACY-" + hashlib.sha1(key.encode()).hexdigest()[:16].upper()


def state_names():
    rows = json.loads((NORMALIZED / "states.json").read_text(encoding="utf-8"))
    return {row["name"].upper(): row["abbr"] for row in rows}


def parse(url, office, abbreviations):
    payload = urlopen(Request(url, headers={"User-Agent": "us-government-data-atlas/0.1"}), timeout=120).read()
    sheet = load_workbook(filename=io.BytesIO(payload), read_only=True, data_only=True).active
    rows = []
    state = ""
    district = "S" if office == "senate" else ""
    for values in sheet.iter_rows(values_only=True):
        cells = list(values)
        texts = [clean(v) for v in cells]
        integer_values = [v for v in cells if isinstance(v, int) and not isinstance(v, bool)]
        if office == "senate":
            first = texts[0] if texts else ""
            state = abbreviations.get(first.split()[0].upper(), state) if first else state
            name = next((text for text in texts if "," in text and "votes" not in text.lower()), "")
            party = next((text for text in texts if text in {"D", "R", "I", "L", "G", "LBT", "LIB", "CON", "CST", "IND", "DFL", "GRN", "W"}), "")
            votes = integer_values[-1] if integer_values else None
        else:
            marker = texts[2] if len(texts) > 2 else ""
            marker_first = marker.splitlines()[0].strip().upper() if marker else ""
            if marker_first in abbreviations:
                state = abbreviations[marker_first]
            elif marker.upper().startswith("DISTRICT "):
                district = marker.splitlines()[0].split()[-1]
                name = marker.splitlines()[1].strip() if "\n" in marker else ""
            name = marker
            party = texts[21] if len(texts) > 21 else ""
            votes = integer_values[-1] if integer_values else None
        if not name or name.lower() in {"scattered", "candidate name"} or "votes:" in name.lower() or name.upper().startswith("2002 "):
            continue
        if not state or not party or party in {"W", ""} or not isinstance(votes, (int, float)):
            continue
        name = name.replace("(I)", "").strip(" ,")
        if not name:
            continue
        rows.append({
            "fec_id": synthetic_id(2002, office, state, district, name, party),
            "candidate_name": name,
            "state_abbreviation": state,
            "d": district,
            "party": party,
            "general_votes": int(votes),
            "election_year": 2002,
            "office": office,
            "source": url,
            "source_sheet": sheet.title,
            "id_type": "synthetic legacy record ID",
        })
    return rows


def main():
    RAW.mkdir(parents=True, exist_ok=True)
    NORMALIZED.mkdir(parents=True, exist_ok=True)
    abbreviations = state_names()
    rows = []
    manifests = {}
    for office, url in SOURCES.items():
        parsed = parse(url, office, abbreviations)
        rows.extend(parsed)
        manifests[office] = {"source": url, "records": len(parsed)}
    rows.sort(key=lambda row: (row["office"], row["state_abbreviation"], row["d"], row["candidate_name"]))
    filename = "fec_federal_elections_2002.json"
    (NORMALIZED / filename).write_text(json.dumps(rows, indent=2) + "\n", encoding="utf-8")
    manifest = {"retrieved_at": datetime.now(timezone.utc).isoformat(), "sources": manifests, "format": "legacy XLSX workbooks", "records": len(rows), "file": "data/normalized/" + filename, "note": "Official FEC compilation; candidate IDs were not present in the 2002 workbooks, so stable synthetic IDs are provided."}
    (RAW / "fec_federal_elections_2002_manifest.json").write_text(json.dumps(manifest, indent=2) + "\n", encoding="utf-8")
    print(json.dumps(manifest, indent=2))


if __name__ == "__main__":
    main()
