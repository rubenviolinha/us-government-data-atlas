#!/usr/bin/env python3
"""Fetch official FEC 2020 presidential general results without an API key."""

from datetime import datetime, timezone
import io, json
from pathlib import Path
from urllib.request import Request, urlopen
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
RAW = ROOT / "data" / "raw"
NORMALIZED = ROOT / "data" / "normalized"
URL = "https://www.fec.gov/documents/4228/federalelections2020.xlsx"


def main():
    RAW.mkdir(parents=True, exist_ok=True); NORMALIZED.mkdir(parents=True, exist_ok=True)
    payload = urlopen(Request(URL, headers={"User-Agent": "us-government-data-atlas/0.1"}), timeout=180).read()
    sheet = load_workbook(filename=io.BytesIO(payload), read_only=True, data_only=True)["9. 2020 Pres General Results"]
    headers = [str(v).strip().lower().replace(" ", "_").replace("#", "") if v is not None else "" for v in next(sheet.iter_rows(values_only=True))]
    rows = []
    for values in sheet.iter_rows(min_row=3, values_only=True):
        row = {h: v for h, v in zip(headers, values) if h and v not in (None, "")}
        if not row.get("fec_id") or not row.get("last_name,__first") or not isinstance(row.get("general_results"), (int, float)):
            continue
        rows.append({"fec_id": row["fec_id"], "candidate_name": row["last_name,__first"], "first_name": row.get("first_name"), "last_name": row.get("last_name"), "state": row.get("state_abbreviation"), "party": row.get("party"), "votes": int(row["general_results"]), "vote_share": row.get("general_%"), "winner": row.get("winner_indicator"), "electoral_votes": row.get("electoral_votes"), "election_year": 2020, "source": URL, "source_sheet": sheet.title})
    rows.sort(key=lambda row: (row["state"], -row["votes"], row["candidate_name"]))
    filename = "fec_presidential_general_2020.json"
    (NORMALIZED / filename).write_text(json.dumps(rows, indent=2) + "\n", encoding="utf-8")
    manifest = {"retrieved_at": datetime.now(timezone.utc).isoformat(), "source": URL, "format": "XLSX workbook", "records": len(rows), "file": "data/normalized/" + filename}
    (RAW / "fec_presidential_general_2020_manifest.json").write_text(json.dumps(manifest, indent=2) + "\n", encoding="utf-8")
    print(json.dumps(manifest, indent=2))


if __name__ == "__main__":
    main()
