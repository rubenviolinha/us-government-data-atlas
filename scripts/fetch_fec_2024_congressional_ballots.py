#!/usr/bin/env python3
"""Fetch the FEC's official 2024 congressional general-ballot candidate workbook."""

from datetime import datetime, timezone
import io
import json
from pathlib import Path
from urllib.request import Request, urlopen

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
NORMALIZED = ROOT / "data" / "normalized"
RAW = ROOT / "data" / "raw"
URL = "https://www.fec.gov/documents/5548/2024congressgecands.xlsx"

def clean(value):
    return value.isoformat() if hasattr(value, "isoformat") else value

def main():
    NORMALIZED.mkdir(parents=True, exist_ok=True)
    RAW.mkdir(parents=True, exist_ok=True)
    payload = urlopen(Request(URL, headers={"User-Agent": "us-government-data-atlas/0.1"}), timeout=300).read()
    (RAW / "fec_2024_congressional_ballots.xlsx").write_bytes(payload)
    workbook = load_workbook(io.BytesIO(payload), read_only=True, data_only=True)
    rows = []
    sheet_counts = {}
    for sheet in workbook.worksheets:
        values = list(sheet.iter_rows(values_only=True))
        if not values:
            continue
        headers = [str(value).strip().lower().replace(" ", "_").replace("#", "number") if value is not None else "" for value in values[0]]
        count = 0
        for values_row in values[1:]:
            row = {header: clean(value) for header, value in zip(headers, values_row) if header and value is not None}
            fec_id = str(row.get("fec_idnumber", row.get("fec_id_number", ""))).strip()
            candidate_name = str(row.get("candidate_name", "")).strip()
            if not fec_id or not candidate_name or fec_id.lower() == "none":
                continue
            district = str(row.get("district", "")).strip()
            office = "senate" if district.startswith("S") else "house"
            row.update({"election_year": 2024, "office": office, "source": URL, "source_sheet": sheet.title})
            row["fec_id_number"] = fec_id
            rows.append(row)
            count += 1
        sheet_counts[sheet.title] = count
    rows.sort(key=lambda row: (row["office"], row.get("state_abbreviation", ""), str(row.get("district", "")), row.get("fec_id_number", ""), row.get("candidate_name", "")))
    filename = "fec_congressional_ballot_candidates_2024.json"
    (NORMALIZED / filename).write_text(json.dumps(rows, indent=2) + "\n", encoding="utf-8")
    manifest = {"retrieved_at": datetime.now(timezone.utc).isoformat(), "source": URL, "format": "Excel workbook", "sheets": sheet_counts, "records": len(rows), "file": "data/normalized/" + filename, "note": "Official FEC compilation of candidates on 2024 congressional general-election ballots; no API key used."}
    (RAW / "fec_2024_congressional_ballots_manifest.json").write_text(json.dumps(manifest, indent=2) + "\n", encoding="utf-8")
    print(json.dumps(manifest, indent=2))

if __name__ == "__main__":
    main()
