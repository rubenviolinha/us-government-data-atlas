#!/usr/bin/env python3
"""Fetch official FEC 2024 presidential general-election results workbook."""

from datetime import datetime, timezone
import json
from pathlib import Path
from urllib.request import Request, urlopen
from zipfile import ZipFile
import io

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
NORMALIZED = ROOT / "data" / "normalized"
RAW = ROOT / "data" / "raw"
URL = "https://www.fec.gov/documents/5645/2024presgeresults.xlsx"


def main():
    NORMALIZED.mkdir(parents=True, exist_ok=True)
    RAW.mkdir(parents=True, exist_ok=True)
    request = Request(URL, headers={"User-Agent": "us-government-data-atlas/0.1"})
    with urlopen(request, timeout=300) as response:
        payload = response.read()
    workbook = load_workbook(io.BytesIO(payload), read_only=True, data_only=True)
    rows = []
    for sheet in workbook.worksheets:
        values = list(sheet.iter_rows(values_only=True))
        if not values:
            continue
        headers = [str(value).strip().lower().replace(" ", "_") if value is not None else "" for value in values[0]]
        for values_row in values[1:]:
            if not any(value is not None for value in values_row):
                continue
            row = {header: value for header, value in zip(headers, values_row) if header}
            row.update({"election_year": 2024, "source": URL, "source_sheet": sheet.title})
            rows.append(row)
    (NORMALIZED / "fec_presidential_general_results_2024.json").write_text(json.dumps(rows, indent=2, default=str) + "\n", encoding="utf-8")
    manifest = {
        "retrieved_at": datetime.now(timezone.utc).isoformat(),
        "source": URL,
        "format": "Excel workbook",
        "sheets": workbook.sheetnames,
        "records": len(rows),
        "file": "data/normalized/fec_presidential_general_results_2024.json",
        "note": "Official results compiled from state and territory election offices; no API key used."
    }
    (RAW / "fec_presidential_general_results_2024_manifest.json").write_text(json.dumps(manifest, indent=2) + "\n", encoding="utf-8")
    print(json.dumps(manifest, indent=2))


if __name__ == "__main__":
    main()
