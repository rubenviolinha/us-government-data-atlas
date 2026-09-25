#!/usr/bin/env python3
"""Fetch official FEC 1996 presidential general results."""

from datetime import datetime, timezone
import hashlib, io, json
from pathlib import Path
from urllib.request import Request, urlopen
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]; RAW = ROOT / "data" / "raw"; NORMALIZED = ROOT / "data" / "normalized"
URL = "https://www.fec.gov/documents/1721/FederalElections96_PresidentialGeneralElectionResults.xlsx"


def main():
    RAW.mkdir(parents=True, exist_ok=True); NORMALIZED.mkdir(parents=True, exist_ok=True)
    sheet = load_workbook(filename=io.BytesIO(urlopen(Request(URL, headers={"User-Agent": "us-government-data-atlas/0.1"}), timeout=120).read()), read_only=True, data_only=True).active
    state = None; rows = []
    for values in sheet.iter_rows(min_row=3, values_only=True):
        text = [str(v).strip() if v is not None else "" for v in values]
        markers = [v for v in text if v and v == v.upper() and v not in {"VOTES:", "TOTAL"} and "," not in v]
        if markers and not any("," in v for v in text):
            state = markers[0]; continue
        name = next((v for v in text if "," in v), "")
        party = next((v for v in text if v in {"D", "R", "I", "LBT", "NL", "W", "AIP", "LIB", "GRE"}), "")
        integers = [v for v in values if isinstance(v, int) and not isinstance(v, bool)]
        votes = integers[-1] if integers else None
        if not state or not name or not party or votes is None or name.lower() in {"scattered", "total"}:
            continue
        key = f"1996|{state}|{name}|{party}"
        rows.append({"fec_id": "LEGACY-" + hashlib.sha1(key.encode()).hexdigest()[:16].upper(), "candidate_name": name, "state": state, "party": party, "votes": int(votes), "election_year": 1996, "source": URL, "source_sheet": sheet.title, "id_type": "synthetic legacy record ID"})
    rows.sort(key=lambda row: (row["state"], -row["votes"], row["candidate_name"]))
    filename = "fec_presidential_general_1996.json"; (NORMALIZED / filename).write_text(json.dumps(rows, indent=2) + "\n", encoding="utf-8")
    manifest = {"retrieved_at": datetime.now(timezone.utc).isoformat(), "source": URL, "format": "XLSX workbook", "records": len(rows), "file": "data/normalized/" + filename, "note": "The historical workbook has no modern FEC candidate IDs; stable synthetic IDs are included."}
    (RAW / "fec_presidential_general_1996_manifest.json").write_text(json.dumps(manifest, indent=2) + "\n", encoding="utf-8"); print(json.dumps(manifest, indent=2))


if __name__ == "__main__": main()
