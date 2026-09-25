#!/usr/bin/env python3
"""Fetch official FEC 2018 and 2016 Senate and House result workbooks."""

from datetime import datetime, timezone
import io
import json
from pathlib import Path
from urllib.request import Request, urlopen

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
NORMALIZED = ROOT / "data" / "normalized"
RAW = ROOT / "data" / "raw"
PUBLICATIONS = {
    2018: "https://www.fec.gov/documents/2706/federalelections2018.xlsx",
    2016: "https://www.fec.gov/documents/1890/federalelections2016.xlsx",
}
SHEETS = ("US Senate Results by State", "US House Results by State")


def clean(value):
    return value.isoformat() if hasattr(value, "isoformat") else value


def fetch(year, url):
    request = Request(url, headers={"User-Agent": "us-government-data-atlas/0.1"})
    with urlopen(request, timeout=300) as response:
        payload = response.read()
    workbook = load_workbook(io.BytesIO(payload), read_only=True, data_only=True)
    rows = []
    sheet_counts = {}
    for sheet_name in SHEETS:
        sheet = next(sheet for sheet in workbook.worksheets if sheet_name in sheet.title)
        values = list(sheet.iter_rows(values_only=True))
        headers = [str(value).strip().lower().replace(" ", "_").replace("#", "") if value is not None else "" for value in values[0]]
        count = 0
        for values_row in values[1:]:
            row = {header: clean(value) for header, value in zip(headers, values_row) if header and value is not None}
            if not row.get("fec_id") or row.get("fec_id") == "n/a" or not row.get("candidate_name"):
                continue
            row.update({"election_year": year, "office": "senate" if "Senate" in sheet_name else "house", "source": url, "source_sheet": sheet.title})
            rows.append(row)
            count += 1
        sheet_counts[sheet_name] = count
    rows.sort(key=lambda row: (row["office"], row.get("state_abbreviation", ""), str(row.get("district", "")), row["fec_id"]))
    filename = f"fec_federal_elections_{year}.json"
    (NORMALIZED / filename).write_text(json.dumps(rows, indent=2) + "\n", encoding="utf-8")
    manifest = {"retrieved_at": datetime.now(timezone.utc).isoformat(), "source": url, "format": "Excel workbook", "sheets": sheet_counts, "records": len(rows), "file": "data/normalized/" + filename, "note": "Official certified federal election compilation; no API key used."}
    (RAW / f"fec_federal_elections_{year}_manifest.json").write_text(json.dumps(manifest, indent=2) + "\n", encoding="utf-8")
    print(json.dumps(manifest, indent=2))


def main():
    NORMALIZED.mkdir(parents=True, exist_ok=True)
    RAW.mkdir(parents=True, exist_ok=True)
    for year, url in PUBLICATIONS.items():
        fetch(year, url)


if __name__ == "__main__":
    main()
